"""Diligencia el formato oficial de la AA3_EV01 (Anexo 1 · Matriz de análisis de riesgo) con el caso de estudio
del punto 6 del componente "Ciberseguridad y seguridad de la información": el proyecto ASEGURA de una unidad
que tramita expedientes en forma presencial y por web (la guía de aprendizaje exige ese caso).

    python generar_formato_oficial.py  → AA3_EV01_Formato_matriz_de_riesgo.xlsx

Los 17 activos y su valor CID salen de la Tabla 4 de la AA1_EV01 (parte 2). Impacto 1-4 a partir del valor CID
(crítico 13-15 = 4, alto 10-12 = 3, medio 7-9 = 2, bajo ≤ 6 = 1). Riesgo = impacto × probabilidad, con los umbrales
de la hoja Fuente del Anexo 1 (medio desde 7, alto desde 10,5). Las fórmulas quedan en Excel.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

AQUI = Path(__file__).parent
SALIDA = AQUI / "AA3_EV01_Formato_matriz_de_riesgo.xlsx"

fino = Side(style="thin", color="808080")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
VERT = Alignment(horizontal="center", vertical="bottom", text_rotation=90, wrap_text=True)
IZQ = Alignment(vertical="center", wrap_text=True)
GRIS, GRIS2, AMARILLO_P, AZUL = "CCCCCC", "E6E6E6", "FFFF99", "1F4E79"
ROJO, NARANJA, VERDE = "FF5B5B", "FFC000", "92D050"
UMBRAL_MEDIO, UMBRAL_ALTO = 7, 10.5

APRENDIZ = [("Aprendiz", "Erick Manuel Sang Cifuentes"),
            ("Programa", "Apropiación de los conceptos en ciberseguridad (23310008) · Ficha 3640188"),
            ("Centro", "Centro de Comercio y Turismo · Regional Quindío"),
            ("Instructora", "Luz Dary Bautista Parra"),
            ("Evidencia", "AA3_EV01 Formato matriz de riesgo · Octubre de 2026")]

# Amenazas del Anexo 1 con la probabilidad estimada para la unidad (1 insignificante … 4 alta)
GRUPOS = [
    ("Amenazas humanas", GRIS, [
        ("Acceso no autorizado ", 4), ("Explotación de errores", 3), ("Ingenieria social", 3), ("Hacker", 3),
        ("Interceptación de datos", 2), ("Codigo malicioso", 3), ("Abuso de la informática", 2), ("Virus", 3)]),
    ("Amenazas del entorno", GRIS2, [
        ("Polución", 1), ("Fallas electricas", 2), ("Sustancias quimicas", 1), ("Fugas de liquido", 1),
        ("Temperatura", 2), ("Incendios", 1), ("Humedad", 1), ("Edificaciones cercanas", 1),
        ("Protección de equipos", 2), ("Fuentes de potencia", 2), ("Cableado datos y electrico", 2),
        ("Ingreso no autorizado", 2), ("Fallas de equipos", 2)]),
    ("Amenzas naturales", AMARILLO_P, [
        ("Terremotos", 2), ("Inundaciones", 1), ("Tormentas electricas", 2), ("Vendabales", 1)]),
    ("Vulnerabilidades ", AMARILLO_P, [
        ("Documentación servidores", 3), ("Manuales y procedimientos", 3), ("Listado de usuarios", 3),
        ("Copias de seguridad", 2), ("Politicas de seguridad", 4), ("Actualizaciones y control de cambios", 3),
        ("Fail over", 3), ("Plan de recuperación de desastres", 4), ("Gestión y monitoreo", 3), ("Auditoria", 3)]),
]
AMENAZAS = [(n, p) for _, _, lista in GRUPOS for n, p in lista]

# Activos de la AA1 (Tabla 4): categoría, activo, C, I, D, ¿maneja expedientes o datos con efectos legales?
ACTIVOS = [
    ("Esencial", "Expedientes en trámite", 5, 5, 4, True),
    ("Esencial", "Expedientes cerrados (archivo central)", 5, 5, 3, True),
    ("Esencial", "Servicio de tramitación electrónica (web)", 4, 5, 5, True),
    ("Esencial", "Servicio de tramitación presencial", 4, 5, 4, True),
    ("Aplicaciones", "Software de tramitación y base de datos", 5, 5, 5, True),
    ("Aplicaciones", "Portal web de administración electrónica", 4, 5, 5, True),
    ("Aplicaciones", "Correo electrónico", 4, 3, 3, False),
    ("Equipamiento", "Servidor", 5, 5, 5, True),
    ("Equipamiento", "Equipos de red (enrutador, cortafuegos)", 4, 5, 5, False),
    ("Equipamiento", "Computadores de escritorio", 4, 4, 3, False),
    ("Equipamiento", "Impresora", 3, 1, 1, False),
    ("Equipamiento", "Fotocopiadora", 1, 1, 1, False),
    ("Comunicaciones", "Conexión a internet", 3, 4, 5, False),
    ("Comunicaciones", "Enlace con el archivo central", 4, 5, 3, True),
    ("Comunicaciones", "Red local (LAN)", 3, 4, 4, False),
    ("Instalaciones", "Edificio de la unidad", 2, 3, 4, False),
    ("Personal", "Funcionario de atención (operador del software)", 3, 4, 4, False),
]


def impacto(c, i, d):
    v = c + i + d
    return 4 if v >= 13 else 3 if v >= 10 else 2 if v >= 7 else 1


# Tratamiento para las amenazas que dejan algún riesgo alto (controles del anexo A de ISO/IEC 27001:2022)
CONTROLES = {
    "Acceso no autorizado ": ("Segundo factor de autenticación en el portal además del NIF y la contraseña; bloqueo tras intentos fallidos; perfiles por rol para el funcionario (A.5.15, A.5.17, A.8.5)", "Servicios informáticos", "2 meses"),
    "Explotación de errores": ("Pruebas de seguridad del portal antes de cada versión y corrección de fallos según OWASP Top 10 (A.8.25, A.8.29)", "Servicios informáticos", "3 meses"),
    "Ingenieria social": ("Sensibilización del personal y avisos a los ciudadanos: la unidad nunca pide la contraseña (A.6.3)", "Atención a usuarios", "2 meses"),
    "Hacker": ("Cortafuegos de aplicaciones web, servidor web separado de la base de datos y pruebas de intrusión anuales (A.8.20, A.8.22)", "Servicios informáticos", "3 meses"),
    "Codigo malicioso": ("Antivirus administrado en servidor y equipos, bloqueo de macros y de dispositivos USB (A.8.7)", "Servicios informáticos", "1 mes"),
    "Virus": ("Antivirus con actualización automática y análisis periódico de los equipos (A.8.7)", "Servicios informáticos", "1 mes"),
    "Documentación servidores": ("Inventario y documentación de servidor, red, software y configuración (A.5.9, A.5.37)", "Servicios informáticos", "2 meses"),
    "Manuales y procedimientos": ("Procedimientos escritos de operación, altas y bajas de usuarios e incidentes (A.5.37)", "Servicios informáticos", "2 meses"),
    "Listado de usuarios": ("Registro de usuarios internos por rol y revisión trimestral de accesos (A.5.18)", "Servicios informáticos", "1 mes"),
    "Politicas de seguridad": ("Política de seguridad de la información aprobada por la dirección y comunicada al personal (A.5.1)", "Dirección y comité de seguimiento", "2 meses"),
    "Actualizaciones y control de cambios": ("Gestión de parches y control de cambios del software de tramitación y del portal (A.8.8, A.8.32)", "Servicios informáticos", "1 mes"),
    "Fail over": ("Servidor o servicio de respaldo y segundo canal de internet para el portal (A.8.14)", "Servicios informáticos", "6 meses"),
    "Plan de recuperación de desastres": ("Plan de continuidad del servicio de tramitación con tiempos objetivo y simulacro anual (A.5.29, A.5.30)", "Comité de seguimiento", "3 meses"),
    "Gestión y monitoreo": ("Registro y revisión de eventos del portal, del servidor y del cortafuegos (A.8.15, A.8.16)", "Servicios informáticos", "2 meses"),
    "Auditoria": ("Auditoría interna semestral del proyecto ASEGURA con lista de chequeo ISO/IEC 27001 (cláusula 9.2)", "Asesoría jurídica y comité", "6 meses"),
}

C0 = 7  # columna G: primera amenaza
F_ENC, F_PROB, F0 = 9, 11, 12  # fila de grupos, fila de probabilidades y primera fila de activos
ULT_COL = C0 + len(AMENAZAS) - 1
ULT_FILA = F0 + len(ACTIVOS) - 1

wb = Workbook()
ws = wb.active
ws.title = "1_Datos"


def celda(f, c, v=None, fill=None, bold=False, al=CENTRO, size=9, hoja=None):
    x = (hoja or ws).cell(f, c, v)
    x.border, x.alignment, x.font = BORDE, al, Font(bold=bold, size=size)
    if fill:
        x.fill = PatternFill("solid", fgColor=fill)
    return x


# Datos del aprendiz y del caso
ws.cell(1, 2, "Matriz de análisis de riesgo · Proyecto ASEGURA: unidad de tramitación de expedientes").font = Font(bold=True, size=13)
ws.cell(2, 2, "Caso de estudio del punto 6 del componente formativo «Ciberseguridad y seguridad de la información» (SENA). "
              "Alcance: tramitación presencial y por web, expedientes, equipos y redes.").font = Font(italic=True, size=9)
for k, (et, val) in enumerate(APRENDIZ):
    celda(3 + k, 2, et, GRIS, True, IZQ)
    ws.merge_cells(start_row=3 + k, start_column=3, end_row=3 + k, end_column=20)
    celda(3 + k, 3, val, al=IZQ)

# Encabezados (estructura del Anexo 1)
ws.merge_cells(start_row=F_ENC, start_column=2, end_row=F_PROB, end_column=2)
celda(F_ENC, 2, "Datos e información\n(activos del caso, valorados en la AA1)", GRIS, True)
ws.merge_cells(start_row=F_ENC, start_column=3, end_row=F_ENC, end_column=5)
celda(F_ENC, 3, "Clasificación", GRIS, True)
for j, h in enumerate(("Confidencial, Privado, Sensitivo", "Obligación por ley / Contrato / Convenio",
                       "Costo de recuperación (tiempo, económico, material, imagen, emocional)")):
    ws.merge_cells(start_row=F_ENC + 1, start_column=3 + j, end_row=F_PROB, end_column=3 + j)
    celda(F_ENC + 1, 3 + j, h, GRIS, True, VERT, 8)
ws.merge_cells(start_row=F_ENC, start_column=6, end_row=F_PROB, end_column=6)
celda(F_ENC, 6, "Impacto\n1 = Insignificante\n2 = Bajo\n3 = Mediano\n4 = Alto", GRIS, True, size=8)
c = C0
for titulo, color, lista in GRUPOS:
    ws.merge_cells(start_row=F_ENC, start_column=c, end_row=F_ENC, end_column=c + len(lista) - 1)
    celda(F_ENC, c, titulo, color, True)
    for n, p in lista:
        celda(F_ENC + 1, c, n.strip(), color, True, VERT, 8)
        celda(F_PROB, c, p, AMARILLO_P, True)
        c += 1
ws.cell(F_ENC - 1, C0, "Probabilidad de la amenaza [1 = Insignificante, 2 = Baja, 3 = Mediana, 4 = Alta] (fila amarilla)").font = Font(italic=True, size=8)

# Activos: clasificación, impacto y riesgo = impacto × probabilidad
for k, (cat, nombre, ci, ii, di, legal) in enumerate(ACTIVOS):
    f = F0 + k
    celda(f, 1, k + 1, size=8)
    celda(f, 2, nombre, al=IZQ)
    for j, marca in enumerate((ci >= 4, legal, ii >= 4 or di >= 4)):
        celda(f, 3 + j, "x" if marca else None)
    celda(f, 6, impacto(ci, ii, di), bold=True)
    for col in range(C0, ULT_COL + 1):
        celda(f, col, f"=$F{f}*{L(col)}${F_PROB}")
rango = f"{L(C0)}{F0}:{L(ULT_COL)}{ULT_FILA}"
for formula, color, op in ((str(UMBRAL_ALTO), ROJO, "greaterThanOrEqual"), (str(UMBRAL_MEDIO), NARANJA, "greaterThanOrEqual"),
                           (str(UMBRAL_MEDIO), VERDE, "lessThan")):
    ws.conditional_formatting.add(rango, CellIsRule(operator=op, formula=[formula],
                                                    fill=PatternFill(start_color=color, end_color=color, fill_type="solid")))

# Leyenda y resumen
f = ULT_FILA + 2
ws.cell(f, 2, "Riesgo = impacto × probabilidad. Umbrales de la hoja Fuente:").font = Font(bold=True, size=9)
formulas = (f'=COUNTIF({rango},">=10.5")', f'=COUNTIF({rango},">=7")-COUNTIF({rango},">=10.5")', f'=COUNTIF({rango},"<7")')
for k, (txt, color) in enumerate((("Alto: 10,5 o más (12 y 16)", ROJO), ("Medio: de 7 a 10,4 (8 y 9)", NARANJA), ("Bajo: menos de 7", VERDE))):
    celda(f + 1 + k, 2, txt, color, al=IZQ)
    ws.merge_cells(start_row=f + 1 + k, start_column=3, end_row=f + 1 + k, end_column=5)
    celda(f + 1 + k, 3, formulas[k])
notas = [
    "Clasificación: «x» en Confidencial si la confidencialidad del activo es 4 o 5; en Obligación por ley si maneja expedientes con efectos "
    "legales o datos personales de los ciudadanos; en Costo de recuperación si su integridad o disponibilidad es 4 o 5 (valuación de la AA1).",
    "Impacto a partir del valor CID de la AA1: crítico (13-15) = 4, alto (10-12) = 3, medio (7-9) = 2, bajo (hasta 6) = 1.",
    "Probabilidad: el portal está expuesto en internet y se accede solo con NIF y contraseña; la unidad aún no tiene política de seguridad "
    "ni plan de recuperación (por eso el proyecto ASEGURA). El detalle de los controles está en la hoja Tratamiento.",
]
for k, n in enumerate(notas):
    ws.cell(f + 5 + k, 2, n).font = Font(italic=True, size=8)

ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 42
for col in "CDE":
    ws.column_dimensions[col].width = 5
ws.column_dimensions["F"].width = 11
for col in range(C0, ULT_COL + 1):
    ws.column_dimensions[L(col)].width = 4.6
ws.row_dimensions[F_ENC + 1].height = 150
ws.row_dimensions[F_ENC].height = 20
for r in range(F0, ULT_FILA + 1):
    ws.row_dimensions[r].height = 24
ws.freeze_panes = ws.cell(F0, C0)
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# --- Hoja Valuación: trazabilidad con la AA1 ---
va = wb.create_sheet("Valuación AA1")
va["A1"] = "Valuación de activos del caso (AA1_EV01, parte 2) y su impacto en la matriz"
va["A1"].font = Font(bold=True, size=13)
for j, h in enumerate(("N.º", "Categoría", "Activo", "C", "I", "D", "Valor CID", "Nivel", "Impacto (1-4)"), 1):
    x = celda(3, j, h, hoja=va)
    x.font, x.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=AZUL)
for k, (cat, nombre, ci, ii, di, _) in enumerate(ACTIVOS):
    v = ci + ii + di
    nivel = "Crítico" if v >= 13 else "Alto" if v >= 10 else "Medio" if v >= 7 else "Bajo"
    for j, val in enumerate((k + 1, cat, nombre, ci, ii, di, v, nivel, impacto(ci, ii, di)), 1):
        celda(4 + k, j, val, hoja=va, al=IZQ if j in (2, 3) else CENTRO)
va.cell(5 + len(ACTIVOS), 1, "Escala CID de 1 (muy bajo) a 5 (muy alto) del material de formación; valor = C + I + D (máximo 15).").font = Font(italic=True, size=8)
for col, w in zip("ABCDEFGHI", (5, 15, 44, 5, 5, 5, 10, 9, 12)):
    va.column_dimensions[col].width = w

# --- Hoja Fuente (igual al Anexo 1) ---
fu = wb.create_sheet("Fuente")
for j, h in enumerate(("Valoración", "Escala", "Valor_min", "Valor_max", "", "Lineas", "Umbral Medio Riesgo", "Umbral Alto Riesgo"), 1):
    fu.cell(1, j, h).font = Font(bold=True)
for k, fila in enumerate((("Ninguna", 1, 1, 3), ("Baja", 2, 4, 6), ("Mediana", 3, 8, 9), ("Alta", 4, 12, 16))):
    for j, v in enumerate(fila, 1):
        fu.cell(2 + k, j, v)
fu.cell(2, 7, UMBRAL_MEDIO), fu.cell(2, 8, UMBRAL_ALTO)
fu.cell(3, 6, "x"), fu.cell(3, 7, "y"), fu.cell(3, 8, "y")
lineas = [1, 1.1, 1.2, 1.3, 1.4, 1.5, 1.6, 1.75, 1.8, 1.9, 2, 2.1, 2.2, 2.3, 2.4, 2.5, 2.625, 2.7, 2.8, 2.9, 3, 3.1, 3.2, 3.3,
          3.4, 3.5, 3.6, 3.7, 3.8, 3.9, 4]
for k, x in enumerate(lineas):
    fu.cell(4 + k, 6, x), fu.cell(4 + k, 7, f"=$G$2/F{4 + k}"), fu.cell(4 + k, 8, f"=$H$2/F{4 + k}")
fu.page_setup.fitToWidth, fu.page_setup.fitToHeight = 1, 1
fu.sheet_properties.pageSetUpPr.fitToPage = True

# --- Hoja Tratamiento: amenazas que dejan riesgo alto ---
tr = wb.create_sheet("Tratamiento")
tr["A1"] = "Tratamiento de los riesgos altos (riesgo ≥ 10,5)"
tr["A1"].font = Font(bold=True, size=13)
tr["A2"] = ("Se agrupan por amenaza o vulnerabilidad: un mismo control reduce el riesgo de todos los activos afectados. "
            "Opción de tratamiento según ISO/IEC 27005: mitigar con controles del anexo A de ISO/IEC 27001:2022.")
tr["A2"].font = Font(italic=True, size=9)
for j, h in enumerate(("Amenaza o vulnerabilidad", "Probabilidad", "Activos con riesgo alto", "Riesgo máximo",
                       "Control propuesto", "Responsable", "Plazo"), 1):
    x = celda(4, j, h, hoja=tr)
    x.font, x.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=AZUL)
fila = 5
for n, p in AMENAZAS:
    afectados = [k + 1 for k, a in enumerate(ACTIVOS) if impacto(*a[2:5]) * p >= UMBRAL_ALTO]
    if not afectados:
        continue
    control, resp, plazo = CONTROLES[n]
    vals = (n.strip(), p, f"{len(afectados)}: n.º " + ", ".join(map(str, afectados)),
            max(impacto(*a[2:5]) for a in ACTIVOS) * p, control, resp, plazo)
    for j, v in enumerate(vals, 1):
        celda(fila, j, v, hoja=tr, al=CENTRO if j in (2, 4, 6, 7) else IZQ)
    tr.cell(fila, 4).fill = PatternFill("solid", fgColor=ROJO)
    tr.row_dimensions[fila].height = 42
    fila += 1
for col, w in zip("ABCDEFG", (26, 11, 24, 9, 66, 16, 9)):
    tr.column_dimensions[col].width = w
tr.page_setup.orientation = "landscape"
tr.page_setup.fitToWidth, tr.page_setup.fitToHeight = 1, 0
tr.sheet_properties.pageSetUpPr.fitToPage = True
tr.cell(fila + 1, 1, "Los números de activo corresponden a la columna N.º de la hoja 1_Datos. Referencias: ISO/IEC 27001:2022; "
                     "ISO/IEC 27005:2022; NIST SP 800-30 Rev. 1 (2012); caso de estudio del componente formativo 1 (SENA).").font = Font(italic=True, size=8)

wb.save(SALIDA)
print("Listo:", SALIDA, "· filas de tratamiento:", fila - 5)

